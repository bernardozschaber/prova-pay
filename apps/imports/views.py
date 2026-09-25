"""Upload -> preview -> confirm flow for payment lists."""
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Prefetch
from django.http import FileResponse, Http404
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST

from openpyxl import load_workbook

from apps.applicators.models import Applicator
from apps.catalog.models import Sector, Unit
from apps.imports.services import (
    all_questions,
    clear_preview,
    confirm_import,
    enrich_preview,
    load_preview,
    stage_uploads,
)
from apps.payroll.models import ImportBatch, ServiceEntry, Shift

ALLOWED_EXTENSIONS = (".xlsx", ".xlsm")


@login_required
def upload(request):
    if request.method == "POST":
        files = [file for file in request.FILES.getlist("files") if file.name.lower().endswith(ALLOWED_EXTENSIONS)]
        if not files:
            messages.error(request, "Selecione ao menos um arquivo .xlsx ou .xlsm.")
            return redirect("imports:upload")
        for error in stage_uploads(request.session, files):
            messages.error(request, f"{error['file_name']}: {error['message']}")
        return redirect("imports:preview")
    return render(request, "imports/upload.html", {"recent_batches": _recent_batches()})


def _recent_batches(limit: int = 10):
    """Os últimos lotes com o que a linha mostra: unidades, total e contagem.

    Os lançamentos já vêm no prefetch, então unidades, soma e contagem saem em
    memória — sem uma consulta por linha da tabela.
    """
    batches = list(
        ImportBatch.objects.select_related("imported_by__profile").prefetch_related(
            Prefetch("entries", queryset=ServiceEntry.objects.select_related("unit__paying_company"))
        )[:limit]
    )
    for batch in batches:
        entries = list(batch.entries.all())
        batch.entry_count = len(entries)
        batch.net_total = sum((entry.net_amount for entry in entries), Decimal("0"))
        codes = {}
        for entry in entries:
            codes.setdefault(entry.unit.short_name, None)
        batch.unit_codes = list(codes)
    return batches


@login_required
def preview(request):
    workbooks = load_preview(request.session)
    if not workbooks:
        return redirect("imports:upload")
    context = {
        "workbooks": enrich_preview(workbooks),
        "questions": all_questions(workbooks),
        # Para a terceira resposta ("associar a outro cadastro"): a busca é no
        # navegador, sobre esta lista, porque o card não recarrega a página.
        "applicators": Applicator.objects.filter(is_active=True).order_by("full_name"),
        "units": Unit.objects.all(),
        "sectors": Sector.objects.all(),
        "shifts": Shift.choices,
        "default_unit": Unit.objects.filter(is_default=True).first(),
        "default_sector": Sector.objects.filter(is_default=True).first(),
        "total_rows": sum(workbook["row_count"] for workbook in workbooks),
    }
    return render(request, "imports/preview.html", context)


@login_required
@require_POST
def confirm(request):
    workbooks = load_preview(request.session)
    if not workbooks:
        return redirect("imports:upload")
    try:
        result = confirm_import(request.POST, workbooks, request.user)
    except ValueError as error:
        messages.error(request, str(error))
        return redirect("imports:preview")
    clear_preview(request.session)
    summary = f"{result['entries']} lançamento(s) importado(s)"
    if result["merged"]:
        summary += f", {result['merged']} lançamento(s) juntado(s) a um cadastro existente"
    if result["applicators"]:
        summary += f", {result['applicators']} cadastro(s) novo(s) criado(s) (primeiro pagamento)"
    if result["skipped"]:
        summary += f", {result['skipped']} linha(s) ignorada(s)"
    messages.success(request, summary + ".")
    if result["duplicates"]:
        # Avulso e em separado: não é detalhe do sucesso, é o operador
        # precisando saber que a lista trouxe serviço que já estava lançado.
        messages.warning(
            request,
            f"{result['duplicates']} linha(s) já estavam lançadas e não entraram de novo — "
            "mesma pessoa, mesma atividade, mesmo dia e mesmo turno. Confira se a data das "
            "abas da planilha está certa.",
        )
    return redirect("payroll:entry_list")


@login_required
@require_POST
def discard(request):
    clear_preview(request.session)
    return redirect("imports:upload")


# --- planilha original -----------------------------------------------------

PREVIEW_MAX_ROWS = 200
PREVIEW_MAX_COLUMNS = 14


def _cell_text(value) -> str:
    """Valor da célula como ela se lê na planilha, não como o Python a imprime."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y") if value.time() == time(0, 0) else value.strftime("%d/%m/%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _read_sheets(batch: ImportBatch) -> list[dict]:
    """Lê a planilha guardada e devolve as abas como linhas de texto."""
    with batch.source_file.open("rb") as handle:
        workbook = load_workbook(BytesIO(handle.read()), read_only=True, data_only=True)
    sheets = []
    try:
        for worksheet in workbook.worksheets:
            rows, truncated = [], False
            for index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if index >= PREVIEW_MAX_ROWS:
                    truncated = True
                    break
                cells = [_cell_text(value) for value in row[:PREVIEW_MAX_COLUMNS]]
                if any(cells):
                    rows.append(cells)
            width = max((len(row) for row in rows), default=0)
            sheets.append({
                "name": worksheet.title,
                "rows": [row + [""] * (width - len(row)) for row in rows],
                "truncated": truncated,
            })
    finally:
        workbook.close()
    return sheets


@login_required
def batch_preview(request, pk: int):
    """Mostra a planilha original do lote, aba por aba, como ela chegou."""
    batch = get_object_or_404(ImportBatch.objects.select_related("imported_by__profile"), pk=pk)
    if not batch.source_file:
        messages.error(request, "Esta importação é anterior ao arquivamento das planilhas; o original não foi guardado.")
        return redirect("imports:upload")
    try:
        sheets = _read_sheets(batch)
    except FileNotFoundError:
        messages.error(request, "A planilha original deste lote não está mais no servidor.")
        return redirect("imports:upload")
    return render(request, "imports/batch_preview.html", {"batch": batch, "sheets": sheets})


def _delete_batches(batches: list[ImportBatch]) -> int:
    """Apaga os lotes com os lançamentos que cada um criou. Devolve o total de lançamentos.

    `ServiceEntry.import_batch` é SET_NULL, então apagar só o lote deixaria os
    lançamentos para trás sem nenhuma origem — e sem forma de encontrá-los pela
    lista que os criou. Os dois saem juntos, na mesma transação, junto com a
    planilha guardada em `media/`, que não serve a mais ninguém depois disso.
    """
    entry_count = ServiceEntry.objects.filter(import_batch__in=batches).count()
    with transaction.atomic():
        ServiceEntry.objects.filter(import_batch__in=batches).delete()
        for batch in batches:
            batch.source_file.delete(save=False)
            batch.delete()
    return entry_count


@login_required
@require_POST
def batch_delete(request, pk: int):
    """Desfaz uma importação: apaga os lançamentos do lote e o próprio lote."""
    batch = get_object_or_404(ImportBatch, pk=pk)
    file_name = batch.file_name
    entry_count = _delete_batches([batch])
    messages.success(request, f"Importação \u201c{file_name}\u201d removida: {entry_count} lançamento(s) excluído(s).")
    return redirect("imports:upload")


@login_required
@require_POST
def batch_bulk_delete(request):
    """O mesmo, para as importações marcadas na tabela.

    Os lotes saem numa transação só: ou a seleção inteira desaparece, ou nada
    desaparece. Meia exclusão deixaria o operador sem saber o que ainda existe.
    """
    # Só dígitos: um "batches" adulterado não pode chegar ao ORM e virar 500.
    selected = [value for value in request.POST.getlist("batches") if value.isdigit()]
    batches = list(ImportBatch.objects.filter(pk__in=selected))
    if not batches:
        messages.error(request, "Selecione ao menos uma importação para excluir.")
        return redirect("imports:upload")
    entry_count = _delete_batches(batches)
    messages.success(
        request,
        f"{len(batches)} importação(ões) removida(s): {entry_count} lançamento(s) excluído(s).",
    )
    return redirect("imports:upload")


@login_required
def batch_download(request, pk: int):
    """Entrega a planilha original. Passa pela sessão: é dado de pagamento."""
    batch = get_object_or_404(ImportBatch, pk=pk)
    if not batch.source_file:
        raise Http404("Planilha original não guardada para esta importação.")
    return FileResponse(batch.source_file.open("rb"), as_attachment=True, filename=batch.file_name)
