"""Lê as duas planilhas do setor que formam o cadastro de aplicadores.

São duas fontes com papéis diferentes:

* a planilha de controle de RPAs, aba "RESUMO DE PGTO POR APLICADOR", diz
  **quem já recebeu** — é ela que define quem entra no cadastro;
* a planilha de agendamento, aba "Aplicadores", diz **o que se sabe** sobre
  cada pessoa (documentos, banco, curso) — é a ficha que preenche o perfil.

A segunda não cria ninguém: só completa quem a primeira já trouxe.
"""
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from apps.applicators.names import is_same_person, normalize_name

PAYMENTS_SHEET = "RESUMO DE PGTO POR APLICADOR"
PAYMENTS_NAME_COLUMN = "APLICADOR"
PROFILES_SHEET = "Aplicadores"
# Os cabeçalhos da ficha, como estão escritos na linha 2 da aba "Aplicadores".
PROFILE_COLUMNS = {
    "nome": "Nome",
    "phone": "Telefone",
    "cpf": "CPF",
    "identity_document": "Identidade",
    "birth_date": "Data Nascimento",
    "gender": "Sexo",
    "neighborhood": "Bairro",
    "vse": "VSE",
    "email": "E-mail",
    "course": "Curso",
    "course_period": "Período",
    "institution": "Instituição",
    "bank_details": "Dados bancários",
    "bank_name": "Banco:",
    "account_type": "Tipo de conta",
    "pix_type": "Tipo PIX",
    "pix_key": "PIX",
    "pis_nit": "PIS / NIT",
    "referral": "Indicação",
}


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for pattern in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(_text(value), pattern).date()
        except ValueError:
            continue
    return None


def read_payment_names(path: str | Path) -> list[str]:
    """Nomes da coluna APLICADOR, na ordem em que aparecem, sem repetir.

    A aba repete o bloco de cabeçalho a cada quinzena, então a leitura é por
    conteúdo da coluna, não por posição: cabeçalho e linha de total saem fora.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[PAYMENTS_SHEET]
        column = None
        names, seen = [], set()
        for row in worksheet.iter_rows(values_only=True):
            cells = [_text(value) for value in row]
            if PAYMENTS_NAME_COLUMN in [cell.upper() for cell in cells]:
                column = [cell.upper() for cell in cells].index(PAYMENTS_NAME_COLUMN)
                continue
            if column is None or column >= len(cells):
                continue
            name = cells[column]
            if not name or name.upper() in {PAYMENTS_NAME_COLUMN, "TOTAL"}:
                continue
            key = normalize_name(name)
            if not key or key in seen:
                continue
            seen.add(key)
            names.append(name)
        if column is None:
            raise ValueError(f'A aba "{PAYMENTS_SHEET}" não tem a coluna "{PAYMENTS_NAME_COLUMN}".')
        return names
    finally:
        workbook.close()
