"""
Generates docs/samples/lista_pagamento_exemplo.xlsx: a payment list in the
"Relatório de Atividade" layout with FICTIONAL applicators, used by
`manage.py seed --demo` and as a reference for the parser.

    .venv/bin/python scripts/make_sample_list.py
"""
from pathlib import Path

from openpyxl import Workbook

OUTPUT = Path(__file__).resolve().parent.parent / "docs" / "samples" / "lista_pagamento_exemplo.xlsx"

# (sheet title, event name, date text, shift, [(name, role, value)])
SHEETS = [
    ("SIMULADO ENEM", "Simulado ENEM II (Prova II)", "12 de Setembro de 2026", "Tarde", [
        ("ALICE FERREIRA DUARTE", "ORIENTADOR", 115), ("BRUNO CASTRO LIMA", "ORIENTADOR", 115),
        ("CAMILA ROCHA NUNES", "APLICADOR", 102), ("DIEGO ALMEIDA PRADO", "APLICADOR", 102),
        ("ELISA MOURA TAVARES (C.E. ONLINE)", "APLICADOR", 115), ("FELIPE GOMES BARBOSA", "APLICADOR", 102),
        ("GABRIELA SANTOS REIS", "APLICADOR", 102), ("HENRIQUE VIEIRA LOPES", "APLICADOR", 102),
    ]),
    ("CMMG", "Simulado CMMG 07", "12 de Setembro de 2026", "Manhã", [
        ("CAMILA ROCHA NUNES", "ORIENTADOR", 115), ("IGOR MARTINS COSTA", "APLICADOR", 102), ("JULIA PEREIRA MENDES", "APLICADOR", 102),
    ]),
    ("ORGANIZAÇÃO ESTOQUE", "Organização do Estoque", "27 de Agosto de 2026", "Tarde", [
        ("DIEGO ALMEIDA PRADO", "APLICADOR", 98), ("JULIA PEREIRA MENDES", "APLICADOR", 98),
    ]),
    ("ATV. PEDAGÓGICA", "Atividade Pedagógica", "05 de Agosto de 2026", "Manhã", [
        ("ALICE FERREIRA DUARTE", "APLICADOR", 98),
    ]),
]


def write_sheet(sheet, event_name, date_text, shift, rows):
    sheet["D1"] = "Relatório de Atividade"
    sheet["D2"], sheet["F2"], sheet["G2"], sheet["H2"] = "Segmento da Atividade:", event_name, "Data:", date_text
    sheet["D3"], sheet["F3"], sheet["G3"], sheet["H3"], sheet["I3"] = "Empresa:", "Matriz", "Valor Total:", f"=SUM(H5:H{4 + len(rows)})", f"Horário: {shift}"
    sheet["D4"], sheet["F4"], sheet["G4"], sheet["H4"], sheet["I4"] = "Nome", "Aplicador", "Orientador (a)", "Valor", "Assinatura"
    for index, (name, role, value) in enumerate(rows, start=5):
        sheet[f"D{index}"] = name
        sheet[f"F{index}" if role == "APLICADOR" else f"G{index}"] = 1
        sheet[f"H{index}"] = value
    total_row = 5 + len(rows) + 1
    sheet[f"F{total_row}"] = f"=SUM(F5:F{total_row - 1})"
    sheet[f"D{total_row + 2}"] = "ASSINATURA DO RESPONSÁVEL PELA CONTRATAÇÃO:"


def main() -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, event_name, date_text, shift, rows in SHEETS:
        write_sheet(workbook.create_sheet(title), event_name, date_text, shift, rows)
    workbook.create_sheet("Planilha1")["A1"] = "aba sem layout, deve ser ignorada"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    print(f"written {OUTPUT}")


if __name__ == "__main__":
    main()
